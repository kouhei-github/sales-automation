"""enriched.json → Excel (シート = 業種)。既存ブック全シートと突き合わせて重複を除外して追記。

使い方:
  python3 export_xlsx.py --in enriched.json --xlsx 営業リスト.xlsx --sheet "塗装業" [--descriptions descriptions.json]
出力: 追加件数 / 重複除外件数 を JSON で stdout に出力
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import dedup_keys, load_json, log, safe_sheet_name, today

# (ヘッダー, enriched キー, 列幅)
COLUMNS = [
    ("会社名", "company", 30),
    ("業種/カテゴリ", "category", 16),
    ("電話番号", "phone", 15),
    ("メールアドレス", "email", 28),
    ("お問い合わせページ", "contact_url", 40),
    ("HP URL", "website", 36),
    ("企業説明(サービス内容)", "description", 60),
    ("代表者名", "representative", 14),
    ("本社住所", "hq_address", 40),
    ("拠点・子会社・支店", "branches", 50),
    ("資本金", "capital", 14),
    ("設立", "established", 12),
    ("従業員数", "employees", 10),
    ("事業内容(HP原文)", "business_raw", 50),
    ("Google Maps評価", "rating", 18),
    ("Google Maps URL", "maps_url", 20),
    ("取得元", "source", 12),
    ("検索クエリ", "query", 18),
    ("取得日", "fetched_at", 11),
]
HEADERS = [c[0] for c in COLUMNS]


def existing_keys(wb):
    keys = set()
    for ws in wb.worksheets:
        hdr = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        if "会社名" not in hdr:
            continue
        idx = {h: i for i, h in enumerate(hdr)}
        for r in ws.iter_rows(min_row=2, values_only=True):
            row = {"会社名": r[idx["会社名"]] if idx.get("会社名") is not None else "",
                   "HP URL": r[idx["HP URL"]] if idx.get("HP URL") is not None else "",
                   "電話番号": r[idx["電話番号"]] if idx.get("電話番号") is not None else ""}
            if row["会社名"]:
                keys |= dedup_keys({k: (v or "") for k, v in row.items()})
    return keys


def style_header(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    for i, (h, _, w) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="inp", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--descriptions", help="{会社名: 説明文} の JSON。あれば description を上書き")
    ap.add_argument("--min-fields", type=int, default=0, help="電話/メール/HP のうち最低いくつ埋まっていれば出力するか")
    ap.add_argument("--limit", type=int, default=0, help="追加する最大件数 (0=無制限)")
    args = ap.parse_args()

    rows = load_json(args.inp, [])
    desc = load_json(args.descriptions, {}) if args.descriptions else {}
    path = Path(args.xlsx)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    known = existing_keys(wb)
    sheet = safe_sheet_name(args.sheet)
    created = sheet not in wb.sheetnames
    ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
    if ws.max_row < 1 or ws.cell(1, 1).value != "会社名":
        style_header(ws)

    added, dup, thin = 0, 0, 0
    for r in rows:
        if args.limit and added >= args.limit:
            break
        keys = dedup_keys(r)
        if keys & known:
            dup += 1
            continue
        filled = sum(1 for k in ("phone", "email", "website") if r.get(k))
        if filled < args.min_fields:
            thin += 1
            continue
        r = dict(r)
        r["fetched_at"] = r.get("fetched_at") or today()
        if r.get("company") in desc and desc[r["company"]]:
            r["description"] = desc[r["company"]]
        vals = [r.get(k, "") or "" for _, k, _ in COLUMNS]
        ws.append(vals)
        n = ws.max_row
        for col, key in ((5, "contact_url"), (6, "website"), (16, "maps_url")):
            v = r.get(key)
            if v and v.startswith("http"):
                ws.cell(n, col).hyperlink = v
                ws.cell(n, col).font = Font(color="0563C1", underline="single")
        for col in (7, 9, 10, 14):
            ws.cell(n, col).alignment = Alignment(wrap_text=True, vertical="top")
        known |= keys
        added += 1
    if created and added == 0:
        wb.remove(ws)  # 追加0件なら空シートを残さない
    if wb.worksheets:
        wb.save(path)
    summary = {"xlsx": str(path.resolve()), "sheet": sheet, "added": added, "duplicates_skipped": dup, "thin_skipped": thin,
               "sheet_total_rows": ws.max_row - 1}
    log("Excel 保存:", json.dumps(summary, ensure_ascii=False))
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
