"""営業リスト作成パイプライン (collect → [descriptions] → export)

  # 1) 収集 + HP巡回 (Google Maps と 求人ボックス の両方、既存Excelと重複する企業は最初から除外)
  python3 run.py collect --query "塗装業" --area "東京都" --target 100 --sheet "塗装業" --xlsx 営業リスト.xlsx
  # 2) (任意) Claude が output/<run>/descriptions.json を書く
  # 3) Excel 出力
  python3 run.py export --run-dir output/塗装業_2026-09-16 --sheet "塗装業" --xlsx 営業リスト.xlsx

collect の要点:
  - --sources gmaps,kb  (既定は両方。gmaps を先に、足りなければ kb)
  - 既存 Excel の全シートを読み、既知企業は候補段階で除外 → 新規 target 件を目指す
  - 途中結果は run-dir に保存 (candidates.json / enriched.json)。再実行時は --resume で続きから
"""
import argparse
import asyncio
import json
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import dedup_keys, load_json, log, safe_sheet_name, save_json, today  # noqa: E402

HERE = Path(__file__).resolve().parent
PY = sys.executable


def sh(*cmd):
    log("$", " ".join(str(c) for c in cmd))
    r = subprocess.run([str(c) for c in cmd], text=True)
    if r.returncode != 0:
        log("!! 失敗 (exit", r.returncode, ")")
    return r.returncode


def known_from_xlsx(xlsx):
    p = Path(xlsx)
    if not p.exists():
        return []
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    rows = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        hdr = next(it, None)
        if not hdr or "会社名" not in hdr:
            continue
        idx = {h: i for i, h in enumerate(hdr)}
        for r in it:
            if r and r[idx["会社名"]]:
                rows.append({"company": r[idx["会社名"]], "website": r[idx.get("HP URL", 0)] or "", "phone": r[idx.get("電話番号", 0)] or ""})
    return rows


def collect(a):
    run_dir = Path(a.run_dir or f"output/{safe_sheet_name(a.sheet or a.query)}_{today()}")
    run_dir.mkdir(parents=True, exist_ok=True)
    known_path = run_dir / "known.json"
    save_json(known_path, known_from_xlsx(a.xlsx))
    log(f"既存Excel既知企業: {len(load_json(known_path, []))} 社 (重複除外に使用)")
    cand = run_dir / "candidates.json"
    if not a.resume and cand.exists():
        cand.unlink()
    sources = [s.strip() for s in a.sources.split(",")]
    need = a.target
    # 求人ボックスは "業" を外したキーワードの方がヒットしやすい (塗装業 → 塗装)
    kb_query = a.kb_query or a.query.rstrip("業")
    for src in sources:
        have = len(load_json(cand, []))
        if have >= a.target:
            break
        want = a.target - have + a.margin
        if src == "gmaps":
            sh(PY, HERE / "search_gmaps.py", "--query", a.query, "--area", a.area, "--limit", want + have, "--out", cand, "--append", "--exclude", known_path)
        elif src == "kb":
            sh(PY, HERE / "search_kyujinbox.py", "--query", kb_query, "--area", a.area, "--limit", want + have, "--out", cand, "--append", "--exclude", known_path)
    cands = load_json(cand, [])
    log(f"候補 {len(cands)} 社 → HP 巡回")
    enriched = run_dir / "enriched.json"
    sh(PY, HERE / "enrich.py", "--in", cand, "--out", enriched, "--concurrency", str(a.concurrency), *(["--resume"] if a.resume else []))
    rows = load_json(enriched, [])
    summary = {
        "run_dir": str(run_dir.resolve()), "candidates": len(cands), "enriched": len(rows), "target": a.target,
        "shortfall": max(0, a.target - len(rows)),
        "with_website": sum(1 for r in rows if r.get("website")), "with_email": sum(1 for r in rows if r.get("email")),
        "with_phone": sum(1 for r in rows if r.get("phone")), "with_representative": sum(1 for r in rows if r.get("representative")),
        "with_capital": sum(1 for r in rows if r.get("capital")),
        "next": "descriptions.json を書いてから `run.py export`" if rows else "候補0件: クエリ/エリアを変えて再実行",
    }
    save_json(run_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=1))


def export(a):
    run_dir = Path(a.run_dir)
    enriched = run_dir / "enriched.json"
    desc = run_dir / "descriptions.json"
    cmd = [PY, HERE / "export_xlsx.py", "--in", enriched, "--xlsx", a.xlsx, "--sheet", a.sheet, "--min-fields", str(a.min_fields), "--limit", str(a.limit)]
    if desc.exists():
        cmd += ["--descriptions", desc]
    else:
        log("descriptions.json なし → 暫定説明 (meta description 等) を使用")
    sys.exit(sh(*cmd))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("collect")
    c.add_argument("--query", required=True)
    c.add_argument("--kb-query", default="", help="求人ボックス用キーワード (既定: query の末尾『業』を除いたもの)")
    c.add_argument("--area", default="")
    c.add_argument("--target", type=int, default=100)
    c.add_argument("--margin", type=int, default=15, help="HP無し等で減る分の上乗せ候補数")
    c.add_argument("--sheet", default="")
    c.add_argument("--xlsx", default="営業リスト.xlsx")
    c.add_argument("--sources", default="gmaps,kb")
    c.add_argument("--run-dir", default="")
    c.add_argument("--concurrency", type=int, default=4)
    c.add_argument("--resume", action="store_true")
    c.set_defaults(fn=collect)
    e = sub.add_parser("export")
    e.add_argument("--run-dir", required=True)
    e.add_argument("--sheet", required=True)
    e.add_argument("--xlsx", default="営業リスト.xlsx")
    e.add_argument("--min-fields", type=int, default=0)
    e.add_argument("--limit", type=int, default=0, help="追加する最大件数 (0=無制限)")
    e.set_defaults(fn=export)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
